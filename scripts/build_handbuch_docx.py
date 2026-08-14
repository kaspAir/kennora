"""Baut das kennora-Benutzerhandbuch als Word-Dokument (docs/kennora-Handbuch.docx).

Nutzt python-docx (auf diesem Rechner verfügbar; kein npm/pandoc/LibreOffice
nötig). Bettet die Screenshots aus docs/bilder/ ein, setzt Überschriften-Stile
(für ein echtes Word-Inhaltsverzeichnis) und die Markenfarben.
Aufruf:  python scripts/build_handbuch_docx.py
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BILD = os.path.join(ROOT, "docs", "bilder")

MARKE = RGBColor(0x40, 0x76, 0xB9)
WARM = RGBColor(0xB0, 0x83, 0x50)     # etwas dunkler als UI-Warm, für Papier lesbar
TINTE = RGBColor(0x23, 0x26, 0x2D)
GRAU = RGBColor(0x6B, 0x70, 0x78)

CENTER = WD_ALIGN_PARAGRAPH.CENTER


# ---- Bausteine --------------------------------------------------------------

def rich(p, teile):
    """Fügt Text mit einfachen Auszeichnungen ein: str | ('b'|'i'|'bi', text)."""
    for teil in teile:
        if isinstance(teil, tuple):
            art, text = teil
            r = p.add_run(text)
            r.bold = "b" in art
            r.italic = "i" in art
        else:
            p.add_run(teil)


def absatz(doc, teile, size=11):
    p = doc.add_paragraph()
    rich(p, teile if isinstance(teile, list) else [teile])
    for r in p.runs:
        r.font.size = Pt(size)
    return p


def liste(doc, punkte):
    for teile in punkte:
        p = doc.add_paragraph(style="List Bullet")
        rich(p, teile if isinstance(teile, list) else [teile])
        for r in p.runs:
            r.font.size = Pt(11)


def bild(doc, name, caption):
    doc.add_picture(os.path.join(BILD, name), width=Inches(6.2))
    doc.paragraphs[-1].alignment = CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = CENTER
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRAU


def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def callout(doc, tag, teile, fill="F0ECE5"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    _shade(cell, fill)
    p0 = cell.paragraphs[0]
    r = p0.add_run(tag.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = GRAU
    p1 = cell.add_paragraph()
    rich(p1, teile)
    for r in p1.runs:
        r.font.size = Pt(11)
    doc.add_paragraph()


def toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for typ, extra in (("begin", None), (None, "instr"), ("separate", None),
                       (None, "text"), ("end", None)):
        if typ:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), typ)
            run._r.append(fc)
        elif extra == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = 'TOC \\o "1-2" \\h \\z \\u'
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Inhaltsverzeichnis – in Word mit F9 aktualisieren."
            run._r.append(t)


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


# ---- Dokument ---------------------------------------------------------------

def main():
    doc = Document()

    # Seitenformat A4 + Ränder
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Mm(22))

    # Grundschriften / Überschriften-Stile
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TINTE
    for lvl, size in ((1, 20), (2, 15)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Georgia"
        st.font.color.rgb = MARKE
        st.font.size = Pt(size)

    # Word soll Felder (TOC) beim Öffnen aktualisieren
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)

    # --- Titelseite ---
    for _ in range(3):
        doc.add_paragraph()
    eb = doc.add_paragraph("B E N U T Z E R H A N D B U C H")
    eb.alignment = CENTER
    for r in eb.runs:
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = WARM
    tit = doc.add_paragraph("kennora")
    tit.alignment = CENTER
    for r in tit.runs:
        r.font.name = "Georgia"
        r.font.size = Pt(48)
        r.font.color.rgb = TINTE
    sub = doc.add_paragraph("Wissen entsteht im Gespräch.")
    sub.alignment = CENTER
    for r in sub.runs:
        r.font.size = Pt(14)
        r.font.color.rgb = GRAU

    toc_break = doc.add_paragraph()
    toc_break.add_run().add_break()  # etwas Luft, dann Seitenumbruch
    doc.add_page_break()

    ov = doc.add_paragraph("Inhalt")
    ov.runs[0].font.name = "Georgia"
    ov.runs[0].font.size = Pt(16)
    ov.runs[0].font.color.rgb = MARKE
    toc(doc)
    doc.add_page_break()

    # --- 01 ---
    h1(doc, "01  Was ist kennora?")
    absatz(doc, ["kennora ist ein ", ("b", "Zuhörer"), ". Du erzählst frei, was du "
                 "weisst – kennora strukturiert das Gesagte sauber, aber menschlich, "
                 "legt es ab und baut über viele Sitzungen hinweg darauf auf."])
    absatz(doc, ["Aus dem, was du sagst, entsteht nach und nach ein ",
                 ("b", "Wissensgraph"), ": deine Gedanken als Knoten, ihre "
                 "Zusammenhänge als Verbindungen. Du siehst ihn auf zwei Arten – als "
                 "aufgeräumten ", ("b", "Baum"), " (wie ein Inhaltsverzeichnis) und "
                 "als assoziatives ", ("b", "Netz"), " (wo Dinge über Themen hinweg "
                 "zusammenklingen)."])
    callout(doc, "Grundhaltung",
            [("b", "Der Mensch führt, nicht die KI. "), "kennora hört zu, ordnet im "
             "Hintergrund und hält sich mit Fragen stark zurück. Es lenkt dich nicht – "
             "du erzählst, kennora bewahrt."])

    # --- 02 ---
    h1(doc, "02  In drei Sätzen")
    liste(doc, [
        [("b", "Reden oder tippen"), " – in jeder Sprache, auch Schweizerdeutsch, "
         "auch gemischt. Kein Formular."],
        ["kennora macht daraus ", ("b", "Aussagen"), ", erkennt den ",
         ("b", "Grundsatz"), " dahinter und ordnet alles in einen wachsenden Baum."],
        ["Du ", ("b", "bestätigst, überarbeitest oder löschst"), " – dein Wissen "
         "bleibt deins."],
    ])

    # --- 03 ---
    h1(doc, "03  Die Startseite")
    absatz(doc, ["Die Startseite ist bewusst leer und ruhig. Ein Klick auf ",
                 ("b", "Sprechen"), " bringt dich in eine Sitzung."])
    bild(doc, "startseite.png", "Die Startseite – das Zeichen, der Name, der Weg hinein.")

    # --- 04 ---
    h1(doc, "04  Eine Sitzung: reden oder tippen")
    absatz(doc, ["Hier passiert alles. Oben stellt kennora eine offene Einladung "
                 "(«Erzähl mir, woran du gerade arbeitest.»). Du kannst ins Textfeld "
                 "tippen oder auf ", ("b", "Sprechen"), " gehen. Mit ",
                 ("b", "Ablegen"), " übergibst du den Beitrag – kennora ordnet ihn "
                 "ein und der Baum darunter wächst."])
    bild(doc, "sitzung.png", "Die Sitzung: oben sprechen oder schreiben, unten dein "
         "Wissen als Baum und die Verbindungen im Netz.")
    absatz(doc, ["Wichtig: Du musst nichts in einem Zug fertig erzählen und nichts "
                 "vollständig sagen. Wissen darf wachsen – beim nächsten Mal knüpft "
                 "kennora an das Gesagte an."])

    # --- 05 ---
    h1(doc, "05  Das Mikrofon (Diktat)")
    absatz(doc, ["Klick auf ", ("b", "Sprechen"), " und rede einfach los. Über das "
                 "kleine Auswahlfeld daneben wählst du dein Mikrofon – aufgenommen "
                 "wird nur dieses, nie das System- oder Meetingaudio. Der erkannte "
                 "Text erscheint im Feld; du liest gegen und legst selbst ab."])
    liste(doc, [
        [("b", "Jede Sprache: "), "kennora erkennt die gesprochene Sprache "
         "automatisch – auch Schweizerdeutsch."],
        [("b", "Sichere Verbindung nötig: "), "Browser geben das Mikrofon nur über "
         "https frei. Auf einer unsicheren Adresse ist der Knopf deaktiviert und "
         "kennora sagt es dir."],
    ])

    # --- 06 ---
    h1(doc, "06  Wie kennora dein Wissen ordnet")
    absatz(doc, ["Jeder Gedanke wird zu einer ", ("b", "Aussage"), ". Eine Aussage "
                 "hat mehrere Ebenen:"])
    liste(doc, [
        [("b", "Kernsatz"), " – sauber formuliert, in deiner Sprache."],
        [("b", "Originalton"), " – deine wörtlichen Worte, bleiben erhalten."],
        [("b", "Grundsatz"), " – das Prinzip dahinter (z. B. «Sorgfalt vor Tempo.»)."],
    ])
    h2(doc, "Reifegrad – wie sicher ist eine Aussage?")
    absatz(doc, ["Vor jeder Aussage steht ein kleines Zeichen. Es zeigt, wie «fest» "
                 "der Gedanke ist:"])
    tab = doc.add_table(rows=3, cols=2)
    tab.style = "Table Grid"
    tab.columns[0].width = Mm(24)
    tab.columns[1].width = Mm(142)
    daten = [("○", "hingeworfen", "frisch von kennora erfasst, noch nicht von dir geprüft."),
             ("✓", "bestätigt", "du hast genickt, der Gedanke stimmt so."),
             ("✎", "überarbeitet", "du hast ihn selbst richtiggestellt; deine Fassung gilt.")]
    for i, (m, name, txt) in enumerate(daten):
        c0 = tab.cell(i, 0)
        c0.width = Mm(24)
        c0.paragraphs[0].alignment = CENTER
        r = c0.paragraphs[0].add_run(m)
        r.font.size = Pt(14)
        c1 = tab.cell(i, 1)
        c1.width = Mm(142)
        rich(c1.paragraphs[0], [("b", name + " – "), txt])
    doc.add_paragraph()
    h2(doc, "Zwei Sichten auf dasselbe Wissen")
    absatz(doc, [("b", "Baum: "), "ein aufgeräumtes Gerüst aus Themen (z. B. "
                 "Schreinerei, Weitergabe), darunter deine konkreten Aussagen – "
                 "übergabetauglich, wie ein Inhaltsverzeichnis."])
    absatz(doc, [("b", "Netz: "), "die nicht-hierarchischen Verbindungen quer über "
                 "die Themen. Hier zeigt kennora, was zusammenklingt – auch über "
                 "Sprachgrenzen hinweg."])
    absatz(doc, ["Verbindungstypen: ", ("b", "gehört-zu, führt-zu, widerspricht, "
                 "reimt-sich-auf, unterscheidet-sich-durch"), "."])
    absatz(doc, [("i", "reimt-sich-auf"), " ist die seltene, überraschende Verbindung "
                 "– zwei Gedanken, die dasselbe Prinzip teilen, obwohl sie aus ganz "
                 "verschiedenen Ecken kommen. Genau diese Funken sind der Kern von "
                 "kennora."])

    # --- 07 ---
    h1(doc, "07  kennora antwortet dir")
    absatz(doc, ["Nach jedem Ablegen spiegelt kennora kurz und warm zurück, was "
                 "angekommen ist und wo etwas gewachsen ist:"])
    callout(doc, "Beispiel einer Rückgabe",
            ["Angekommen ist deine Sorgfalt im Handwerk – und dass sie für dich mit "
             "dem Weitergeben zusammenhängt. Dein Bereich Schreinerei nimmt Form an."],
            fill="F4ECE2")
    absatz(doc, [("b", "Und die Fragen? "), "kennora fragt selten – nie nach jedem "
                 "Beitrag. Nur wenn eine echte Lücke offen bleibt, kommt höchstens "
                 "eine zurückhaltende Frage. Das ist Absicht: Du führst das Gespräch, "
                 "nicht kennora."])

    # --- 08 ---
    h1(doc, "08  Dein Wissen kuratieren")
    absatz(doc, ["Neben jeder Aussage im Baum findest du drei Aktionen:"])
    liste(doc, [
        [("b", "Bestätigen (✓)"), " – dein Nicken. Die Aussage wird bestätigt."],
        [("b", "Überarbeiten (✎)"), " – öffnet eine Seite, auf der du Kernsatz und "
         "Grundsatz korrigierst. Dein Originalton bleibt erhalten, und deine Fassung "
         "wird die massgebliche."],
        [("b", "Löschen (🗑)"), " – entfernt die Aussage (mit Rückfrage)."],
    ])
    bild(doc, "bearbeiten.png", "Überarbeiten: du stellst richtig, was kennora falsch "
         "verstanden hat – der Originalton bleibt als Referenz sichtbar.")
    callout(doc, "Warum das wichtig ist",
            ["kennora kann eine Nuance verdrehen (etwa Ursache und Wirkung). Beim "
             "Gegenlesen fängst du das – und deine Korrektur zählt."])

    # --- 09 ---
    h1(doc, "09  In jeder Sprache")
    absatz(doc, ["Rede, wie dir der Schnabel gewachsen ist – Hochdeutsch, "
                 "Schweizerdeutsch, Französisch, Englisch, auch gemischt. kennora "
                 "antwortet in deiner Sprache und formuliert den Kernsatz in "
                 "derselben Sprache, in der du gesprochen hast."])
    absatz(doc, ["Das Besondere: Verbindungen entstehen über Sprachgrenzen hinweg – "
                 "ein französischer Gedanke kann sich auf einen deutschen «reimen», "
                 "weil beide dasselbe Prinzip teilen. Bei Deutsch schreibt kennora "
                 "immer in Schweizer Rechtschreibung («ss» statt «ß»)."])

    # --- 10 ---
    h1(doc, "10  Testdaten zurücksetzen")
    absatz(doc, ["Zum Ausprobieren gibt es auf der Entwicklungs-Umgebung den roten "
                 "Knopf «Testdaten zurücksetzen». Er löscht alle Aussagen und gibt "
                 "dir ein sauberes Blatt. Diese Funktion existiert bewusst nur auf "
                 "dev – in der richtigen Nutzung ist sie nicht vorhanden."])

    # --- 11 ---
    h1(doc, "11  Tipps")
    liste(doc, [
        [("b", "Einfach erzählen."), " Nicht auf Vollständigkeit oder schöne Sätze "
         "achten – kennora glättet für dich."],
        [("b", "Wachsen lassen."), " Mehrere kurze Sitzungen sind besser als eine "
         "perfekte. kennora knüpft immer an."],
        [("b", "Gegenlesen."), " Überflieg die Rückgabe und den Baum; wo etwas nicht "
         "stimmt, ein Klick auf Überarbeiten und richtigstellen."],
        [("b", "Grundsätze pflegen."), " Die Prinzipien sind das Herz – je klarer "
         "sie sind, desto bessere Verbindungen findet kennora."],
    ])

    schluss = doc.add_paragraph("kennora — Wissen entsteht im Gespräch.")
    schluss.alignment = CENTER
    for r in schluss.runs:
        r.font.color.rgb = GRAU
        r.font.size = Pt(10)

    ziel = os.path.join(ROOT, "docs", "kennora-Handbuch.docx")
    doc.save(ziel)
    print("Word-Handbuch geschrieben:", os.path.normpath(ziel))


if __name__ == "__main__":
    main()
